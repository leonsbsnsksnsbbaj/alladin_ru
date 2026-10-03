"""Сверка Python-калькулятора с Excel-книгой.

1) Базовый сценарий: все ячейки, сохранённые Excel (ежедневная таблица и итоги).
2) Матрица сценариев: копии книги с другими входами пересчитываются
   LibreOffice (headless) и сравниваются с Python.
Запуск: python3 verify.py путь/к/книге.xlsx
"""
import itertools, shutil, subprocess, sys, tempfile
from datetime import date, datetime
from pathlib import Path

import openpyxl
from spectranshi import Inputs, calculate

TOL = 0.005  # полкопейки
MAIN, CALC = "Главная страница", "Расчет"


def as_date(v):
    return v.date() if isinstance(v, datetime) else v


def compare(ws_calc, ws_main, ws_solver, r, label):
    """Возвращает (кол-во проверенных значений, макс. отклонение в рублях, макс. в ставках, ошибки)."""
    errs, n, max_rub, max_rate = [], 0, 0.0, 0.0

    def chk(name, excel, py, money):
        nonlocal n, max_rub, max_rate
        n += 1
        if isinstance(excel, (datetime, date)):
            if as_date(excel) != py:
                errs.append(f"{label} {name}: excel={excel} py={py}")
            return
        d = abs(float(excel or 0) - float(py))
        if money:
            max_rub = max(max_rub, d)
            if d >= TOL:
                errs.append(f"{label} {name}: excel={excel} py={py} diff={d}")
        else:
            max_rate = max(max_rate, d)
            if d > 1e-12:
                errs.append(f"{label} {name}: excel={excel} py={py} diff={d}")

    chk("Расчет!C4", ws_calc["C4"].value, r.end_date, False)
    chk("Расчет!C5", ws_calc["C5"].value, r.amount, True)
    chk("Расчет!C7", ws_calc["C7"].value, r.placement_rate, False)
    chk("Расчет!C11", ws_calc["C11"].value, r.mf, False)
    chk("Расчет!C14", ws_calc["C14"].value, r.simple_yield, False)
    chk("Расчет!C16", ws_calc["C16"].value, r.gross_income, True)
    chk("Расчет!C18", ws_calc["C18"].value, r.final_payout, True)
    chk("Расчет!C19", ws_calc["C19"].value, r.final_yield, False)
    chk("Расчет!H4", ws_calc["H4"].value, r.mf_total, True)
    chk("Решатель!C16", ws_solver["C16"].value, r.solver_exact[6], False)
    chk("Решатель!C17", ws_solver["C17"].value, r.solver_rates[6], False)
    chk("Решатель!C25", ws_solver["C25"].value, r.solver_exact[9], False)
    chk("Решатель!C26", ws_solver["C26"].value, r.solver_rates[9], False)
    money_keys = {"Доход клиента гросс", "Доход минус налог", "Доход итоговый* минус налог",
                  "СС по Дожитию", "СС по СНС"}
    cells = ["D10", "D17", "D18", "D19", "D20", "D21", "D22", "D24", "D25", "D26", "D27", "D28", "D29"]
    keys = [k for k in r.main_page if k != "Доход клиента гросс"] 
    keys.insert(3, "Доход клиента гросс")
    for cell, key in zip(cells, keys):
        chk(f"Главная!{cell}", ws_main[cell].value, r.main_page[key], key in money_keys)
    # ежедневная таблица
    excel_rows = 0
    for i in range(4, 623):
        if ws_calc[f"E{i}"].value in (None, ""):
            break
        excel_rows += 1
        row = r.rows[i - 4]
        chk(f"Расчет!E{i}", ws_calc[f"E{i}"].value, row.day, False)
        chk(f"Расчет!F{i}", ws_calc[f"F{i}"].value, row.deposit, True)
        chk(f"Расчет!G{i}", ws_calc[f"G{i}"].value, row.mf_day, True)
        chk(f"Расчет!J{i}", ws_calc[f"J{i}"].value, row.deposit_net, True)
    if excel_rows != len(r.rows):
        errs.append(f"{label}: строк в таблице excel={excel_rows} py={len(r.rows)}")
    return n, max_rub, max_rate, errs


def check_saved(src):
    wb = openpyxl.load_workbook(src, data_only=True)
    start = as_date(wb[CALC]["E4"].value)
    r = calculate(Inputs(term_months=wb[MAIN]["D5"].value, premium=wb[MAIN]["D6"].value,
                         ukv=wb[MAIN]["D8"].value, start_date=start))
    return compare(wb[CALC], wb[MAIN], wb["Решатель"], r, f"[сохранённый Excel, старт {start}]")


SCENARIOS = list(itertools.product(
    [1, 2, 3, 4, 6, 9, 12],                                          # срок
    [50_000_000, 100_000_000, 777_777_777.77, 20_000_000_000],       # премия
    [0.0013, 0.0033],                                                # УКВ
    [date(2026, 9, 26), date(2027, 1, 31), date(2024, 8, 31), date(2027, 12, 30)],  # старт
))


def check_libreoffice(src):
    tmp = Path(tempfile.mkdtemp())
    (tmp / "out").mkdir()
    for k, (term, prem, ukv, start) in enumerate(SCENARIOS):
        wb = openpyxl.load_workbook(src)
        wb[MAIN]["D5"], wb[MAIN]["D6"], wb[MAIN]["D8"] = term, prem, ukv
        wb[CALC]["E4"] = datetime(start.year, start.month, start.day)  # вместо TODAY()+1
        wb.save(tmp / f"s{k:03d}.xlsx")
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir",
                    str(tmp / "out"), *map(str, sorted(tmp.glob("s*.xlsx")))],
                   check=True, capture_output=True)
    total = (0, 0.0, 0.0, [])
    for k, (term, prem, ukv, start) in enumerate(SCENARIOS):
        wb = openpyxl.load_workbook(tmp / "out" / f"s{k:03d}.xlsx", data_only=True)
        r = calculate(Inputs(term_months=term, premium=prem, ukv=ukv, start_date=start))
        res = compare(wb[CALC], wb[MAIN], wb["Решатель"], r,
                      f"[LO: {term} мес, {prem:,.2f}, УКВ {ukv}, старт {start}]")
        total = (total[0] + res[0], max(total[1], res[1]), max(total[2], res[2]), total[3] + res[3])
    shutil.rmtree(tmp)
    return total


def report(title, res):
    n, rub, rate, errs = res
    print(f"{title}: проверено {n} значений, макс. расхождение {rub:.3e} руб., "
          f"{rate:.3e} по ставкам — {'OK' if not errs else f'ОШИБОК: {len(errs)}'}")
    for e in errs[:20]:
        print("   ", e)
    return not errs


if __name__ == "__main__":
    src = sys.argv[1]
    ok = report("1. Сохранённые значения Excel", check_saved(src))
    ok &= report(f"2. LibreOffice, {len(SCENARIOS)} сценариев", check_libreoffice(src))
    sys.exit(0 if ok else 1)
